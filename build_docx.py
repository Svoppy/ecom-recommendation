from pathlib import Path
import json
import re
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
SRC = (ROOT / "article_body.md").read_text(encoding="utf-8")
RESULT = json.loads((ROOT / "results/model_results.json").read_text(encoding="utf-8"))
TITLE = "Regression Models for Product Recommendation in a Kazakhstan E-Commerce Dataset"
ABSTRACT = (
    "This study compares regression models as scoring functions for an implicit-feedback product recommender. "
    "The experiment uses the Kazakhstan kz.csv extract of the Kaggle eCommerce purchase-history dataset. "
    "After requiring valid 2020 timestamps, product identifiers, prices, and user identifiers, 511,878 of 2,633,521 "
    "source rows remained. Events before 1 September 2020 formed the historical feature window; later events supplied "
    "observed user-product positives. A reproducible sample of 6,000 positive pairs was paired with sampled items "
    "per positive as assumed negatives, producing 29,318 candidate rows for 4,709 users. Fourteen historical and item "
    "attributes were used to score one binary relevance target. Eleven regression estimators were compared using ten-fold "
    "GroupKFold by user. LightGBM achieved the lowest mean RMSE (0.3349 ± 0.0033) and highest mean R² (0.2444 ± 0.0143); "
    "XGBoost was effectively tied on RMSE, while Gradient Boosting Regression attained the highest NDCG@10 (0.7467). "
    "These offline results describe ranking among sampled candidates, not customer satisfaction or causal sales impact. "
    "The high rate of missing user identifiers and the absence of impression logs make the estimates conditional and "
    "call for validation on current, fully observed marketplace traffic."
)
KEYWORDS = "e-commerce; product recommendation; implicit feedback; Kazakhstan; regression; offline evaluation"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)

def margins(cell, top=75, start=75, bottom=75, end=75):
    tcPr = cell._tc.get_or_add_tcPr()
    node = tcPr.first_child_found_in("w:tcMar")
    if node is None:
        node = OxmlElement("w:tcMar"); tcPr.append(node)
    for side, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        item = node.find(qn("w:" + side))
        if item is None:
            item = OxmlElement("w:" + side); node.append(item)
        item.set(qn("w:w"), str(val)); item.set(qn("w:type"), "dxa")

def repeat_header(row):
    prop = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:tblHeader"); tag.set(qn("w:val"), "true"); prop.append(tag)

def chart_file():
    rows = sorted(RESULT["results"], key=lambda x: x["rmse_mean"], reverse=True)
    names = [x["algorithm"] for x in rows]
    vals = [x["rmse_mean"] for x in rows]
    errs = [x["rmse_std"] for x in rows]
    colors = ["#2D6A73" if n in ("LightGBM", "XGBoost", "CatBoost", "HistGradientBoosting") else "#86A9AA" for n in names]
    fig, ax = plt.subplots(figsize=(8, 5.4))
    ax.barh(names, vals, xerr=errs, color=colors, capsize=2)
    ax.set_xlabel("RMSE (lower is better)")
    ax.set_title("Ten-fold grouped validation")
    ax.grid(axis="x", color="#DDE4E5", linewidth=.7); ax.set_axisbelow(True)
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.tick_params(axis="y", length=0, labelsize=8)
    fig.tight_layout()
    out = ROOT / ".build" / "rmse_comparison.png"; out.parent.mkdir(exist_ok=True)
    fig.savefig(out, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return out

def add_results_table(doc):
    doc.add_heading("Table 1. Regression model comparison", level=2)
    heads = ["Algorithm", "Features", "Targets", "k-folds", "RMSE mean ± SD", "R² mean ± SD", "NDCG@10"]
    widths = [1.48, .50, .50, .50, 1.05, .98, .78]
    table = doc.add_table(rows=1, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    for j, label in enumerate(heads):
        c = table.rows[0].cells[j]; c.text = label; c.width = Inches(widths[j])
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        margins(c); shade(c, "244E58")
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in c.paragraphs[0].runs:
            run.font.bold = True; run.font.size = Pt(7.4); run.font.color.rgb = RGBColor(255,255,255)
    repeat_header(table.rows[0])
    for i, row in enumerate(RESULT["results"]):
        values = [row["algorithm"], str(row["features"]), str(row["targets"]), str(row["folds"]),
                  f'{row["rmse_mean"]:.4f} ± {row["rmse_std"]:.4f}',
                  f'{row["r2_mean"]:.4f} ± {row["r2_std"]:.4f}',
                  f'{row["ndcg_at_10_mean"]:.4f}']
        cells = table.add_row().cells
        for j, val in enumerate(values):
            c = cells[j]; c.text = val; c.width = Inches(widths[j])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(c)
            if i % 2: shade(c, "F0F5F5")
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            for run in c.paragraphs[0].runs: run.font.size = Pt(7.3)
    note = doc.add_paragraph("Note. All models used 14 predictors, one binary implicit-relevance target, and ten GroupKFold folds grouped by user. Lower RMSE and higher R²/NDCG@10 are preferred. SD is the standard deviation across folds. Negative labels are sampled non-purchases, not observed dislikes.")
    note.paragraph_format.space_before = Pt(4); note.paragraph_format.space_after = Pt(5)
    for run in note.runs: run.italic = True; run.font.size = Pt(8)
    fig = doc.add_paragraph(); fig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fig.paragraph_format.space_before = Pt(2); fig.paragraph_format.space_after = Pt(0)
    fig.add_run().add_picture(str(chart_file()), width=Inches(5.9))
    cap = doc.add_paragraph("Figure 1. Mean ten-fold RMSE; error bars show one standard deviation.")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs: run.italic = True; run.font.size = Pt(8)

def main():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.72); sec.bottom_margin = Inches(.7)
    sec.left_margin = Inches(.78); sec.right_margin = Inches(.78)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"; normal.font.size = Pt(10); normal.font.color.rgb = RGBColor(24,34,38)
    normal.paragraph_format.line_spacing = 1.08; normal.paragraph_format.space_after = Pt(5)
    for style_name, size in (("Heading 1", 14), ("Heading 2", 11)):
        style = doc.styles[style_name]; style.font.name = "Aptos Display"; style.font.size = Pt(size)
        style.font.bold = True; style.font.color.rgb = RGBColor(0,0,0)
        style.paragraph_format.space_before = Pt(12 if style_name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(4)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(TITLE); run.font.name = "Aptos Display"; run.font.size = Pt(20); run.font.bold = True; run.font.color.rgb = RGBColor(0,0,0)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Dias Kazikhanov  |  Astana IT University  |  2026"); r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(75,91,96)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Abstract"); r.bold = True; r.font.size = Pt(10.5)
    p = doc.add_paragraph(ABSTRACT); p.paragraph_format.line_spacing = 1.02; p.paragraph_format.space_after = Pt(5)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Keywords: "); r.bold = True; p.add_run(KEYWORDS)

    in_refs = False
    for block in re.split(r"\n\s*\n", SRC):
        text = block.strip()
        if not text: continue
        if text == "## References":
            in_refs = True
            doc.add_heading("References", level=1)
            continue
        if text.startswith("#"):
            level = len(text) - len(text.lstrip("#"))
            label = text[level:].strip()
            doc.add_heading(label, level=1 if level == 2 else 2)
            continue
        if in_refs:
            for line in text.splitlines():
                line = line.strip()
                if not line:
                    continue
                p = doc.add_paragraph(line)
                p.paragraph_format.left_indent = Inches(.22); p.paragraph_format.first_line_indent = Inches(-.22)
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs: r.font.size = Pt(8.5)
            continue
        text = text.replace(chr(96), "")
        p = doc.add_paragraph(text); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.widow_control = True
        if text.startswith("Table 1 presents"):
            add_results_table(doc)

    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE"); footer._p.append(fld)
    out = ROOT / "research_paper.docx"; doc.save(out)
    body = SRC.split("## Conclusion", 1)[0]
    body = re.sub(r"^#+.*$", "", body, flags=re.M)
    count = len(re.findall(r"\b[\w–'-]+\b", body))
    print(f"Created {out}\nNarrative body words (excluding headings and requested exclusions): {count}\nReferences: 15")

if __name__ == "__main__": main()
